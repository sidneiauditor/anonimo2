import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

OUTPUT = "Manual_Anonimo_v1.0.docx"
LOGO   = "logo_anonimo.png"

AZUL_ESCURO = RGBColor(0x0F, 0x17, 0x2A)
AZUL        = RGBColor(0x25, 0x63, 0xEB)
AZUL_LINHA  = RGBColor(0x3B, 0x82, 0xF6)
VERDE_ESC   = RGBColor(0x06, 0x4E, 0x3B)
CINZA       = RGBColor(0x64, 0x74, 0x8B)
BRANCO      = RGBColor(0xFF, 0xFF, 0xFF)
AMARELO     = RGBColor(0xFC, 0xD3, 0x4D)

doc = Document()

# ── Margens ───────────────────────────────────────────────────────────────────
for section in doc.sections:
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2.5)

# ── Helpers ───────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, color="E2E8F0"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def para(text, bold=False, size=11, color=None, align=WD_ALIGN_PARAGRAPH.LEFT,
         space_before=0, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p

def heading(text, level=1, size=16, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color or AZUL_ESCURO
    return p

def subheading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = AZUL
    return p

def hr_line(color_hex="E2E8F0"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def caixa(text, bg_hex="FFFBEB", border_hex="FCD34D", bold_prefix=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.columns[0].width = Cm(16)
    cell = t.cell(0, 0)
    set_cell_bg(cell, bg_hex)
    set_cell_border(cell, border_hex)
    cell.width = Cm(16)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    if bold_prefix:
        r = p.add_run(bold_prefix + " ")
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
        r2 = p.add_run(text)
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
    else:
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def caixa_info(text):
    return caixa(text, bg_hex="EFF6FF", border_hex="2563EB")

def page_break():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)

# ── CAPA ──────────────────────────────────────────────────────────────────────
# Fundo escuro via tabela de 1 célula
capa = doc.add_table(rows=1, cols=1)
capa.alignment = WD_TABLE_ALIGNMENT.CENTER
capa.columns[0].width = Cm(16)
capa_cell = capa.cell(0, 0)
set_cell_bg(capa_cell, "0F172A")
capa_cell.width = Cm(16)
capa_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Logo
if os.path.exists(LOGO):
    p_logo = capa_cell.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(30)
    p_logo.paragraph_format.space_after  = Pt(10)
    run_logo = p_logo.add_run()
    run_logo.add_picture(LOGO, width=Cm(4))

# Título
p_titulo = capa_cell.add_paragraph()
p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_titulo.paragraph_format.space_before = Pt(4)
p_titulo.paragraph_format.space_after  = Pt(4)
r = p_titulo.add_run("Anônimo")
r.bold = True
r.font.size = Pt(32)
r.font.color.rgb = BRANCO

# Subtítulo
p_sub = capa_cell.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_before = Pt(0)
p_sub.paragraph_format.space_after  = Pt(4)
r = p_sub.add_run("Anonimização de Dados Sensíveis")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

# Versão
p_ver = capa_cell.add_paragraph()
p_ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ver.paragraph_format.space_before = Pt(0)
p_ver.paragraph_format.space_after  = Pt(20)
r = p_ver.add_run("v1.0  ·  Manual do Utilizador")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x60, 0xA5, 0xFA)

# Linha azul (parágrafo com borda inferior)
p_linha = capa_cell.add_paragraph()
p_linha.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_linha.paragraph_format.space_before = Pt(4)
p_linha.paragraph_format.space_after  = Pt(20)
pPr = p_linha._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '12')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '3B82F6')
pBdr.append(bottom)
pPr.append(pBdr)
# Indentação para simular linha centralizada
p_linha._p.get_or_add_pPr()
ind = OxmlElement('w:ind')
ind.set(qn('w:left'), '1800')
ind.set(qn('w:right'), '1800')
pPr.append(ind)

# Info rodapé da capa
p_coord = capa_cell.add_paragraph()
p_coord.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_coord.paragraph_format.space_before = Pt(4)
p_coord.paragraph_format.space_after  = Pt(2)
r = p_coord.add_run("Coordenadoria de Inteligência Fiscal")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

p_inf = capa_cell.add_paragraph()
p_inf.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inf.paragraph_format.space_before = Pt(0)
p_inf.paragraph_format.space_after  = Pt(30)
r = p_inf.add_run("Processamento 100% local  ·  Sem envio de dados")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

page_break()

# ── ÍNDICE ────────────────────────────────────────────────────────────────────
heading("Índice", size=16)
hr_line()

indice_items = [
    ("1.", "Introdução", "2"),
    ("2.", "Requisitos do Sistema", "2"),
    ("3.", "Como Iniciar o Sistema", "3"),
    ("4.", "Guia de Uso — Passo a Passo", "3"),
    ("    4.1", "Carregar Arquivo", "3"),
    ("    4.2", "Analisar Estrutura", "4"),
    ("    4.3", "Selecionar Campos", "4"),
    ("    4.4", "Anonimizar Dados", "4"),
    ("    4.5", "Reverter Tokens", "5"),
    ("    4.6", "Processar PDF / OCR", "5"),
    ("5.", "Segurança e Privacidade", "5"),
]
for num, titulo, pag in indice_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    tab = p.paragraph_format.tab_stops
    from docx.shared import Pt as PT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r1 = p.add_run(f"{num}  {titulo}")
    r1.font.size = Pt(10)
    r2 = p.add_run(f"\t{pag}")
    r2.font.size  = Pt(10)
    r2.bold = True
    r2.font.color.rgb = AZUL

page_break()

# ── 1. INTRODUÇÃO ──────────────────────────────────────────────────────────────
heading("1. Introdução")
hr_line()
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("O ")
r.font.size = Pt(10)
r = p.add_run("Anônimo v1.0")
r.bold = True; r.font.size = Pt(10)
r = p.add_run(" é uma ferramenta de anonimização de dados sensíveis desenvolvida para uso interno da ")
r.font.size = Pt(10)
r = p.add_run("Coordenadoria de Inteligência Fiscal")
r.bold = True; r.font.size = Pt(10)
r = p.add_run(". Permite substituir informações confidenciais — como CNPJ, CPF, razão social e endereços — "
              "por tokens criptográficos determinísticos, garantindo proteção dos dados sem perda de rastreabilidade.")
r.font.size = Pt(10)

para("Todo o processamento ocorre exclusivamente no computador local. Nenhum dado é transmitido pela internet, "
     "tornando a ferramenta adequada para ambientes com restrições de segurança.", size=10, space_after=6)

caixa_info("Formatos suportados: .xlsx (Excel), .csv, .docx (Word), .txt e .pdf (com OCR)")

# ── 2. REQUISITOS ──────────────────────────────────────────────────────────────
heading("2. Requisitos do Sistema")
hr_line()

req_data = [
    ("Componente", "Requisito", True),
    ("Sistema Operacional", "Windows 10 / 11", False),
    ("Python", "Versão 3.8 ou superior (para servidor local)", False),
    ("Navegador", "Google Chrome, Edge ou Firefox (versão atual)", False),
    ("Memória RAM", "Mínimo 4 GB (recomendado 8 GB para arquivos grandes)", False),
    ("Espaço em disco", "Pasta do sistema + arquivos a processar", False),
    ("Rede", "Não necessária — funcionamento 100% offline", False),
]
t_req = doc.add_table(rows=len(req_data), cols=2)
t_req.alignment = WD_TABLE_ALIGNMENT.LEFT
col_w = [Cm(5), Cm(11)]
for i, (c1, c2, header) in enumerate(req_data):
    row = t_req.rows[i]
    for j, (cell, txt, w) in enumerate(zip(row.cells, [c1, c2], col_w)):
        cell.width = w
        set_cell_border(cell)
        if header:
            set_cell_bg(cell, "0F172A")
        elif i % 2 == 0:
            set_cell_bg(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(txt)
        r.bold = header
        r.font.size = Pt(9)
        r.font.color.rgb = BRANCO if header else RGBColor(0x1E, 0x29, 0x3B)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ── 3. COMO INICIAR ────────────────────────────────────────────────────────────
heading("3. Como Iniciar o Sistema")
hr_line()
para("O Anônimo requer um servidor local para funcionar. Os módulos PDF.js, Tesseract (OCR) e processamento "
     "de planilhas não operam via protocolo file://.", size=10, space_after=6)

subheading("Procedimento de inicialização:")

passos = [
    ("1", "Localize o arquivo Iniciar-Anonimo.bat",
     "Abra o Windows Explorer e navegue até a pasta do sistema."),
    ("2", "Execute com duplo-clique",
     "Clique duas vezes sobre o arquivo. Uma janela de terminal será aberta."),
    ("3", "Aguarde o servidor iniciar",
     "O terminal exibirá a mensagem confirmando que o servidor está rodando na porta 8000."),
    ("4", "O navegador abrirá automaticamente",
     "A aplicação será aberta em http://localhost:8000/Anonimo.html."),
    ("5", "Para encerrar",
     "Clique na janela do terminal e pressione Ctrl+C para parar o servidor."),
]
for num, titulo, desc in passos:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(1)
    p.paragraph_format.left_indent  = Cm(0.5)
    r = p.add_run(f"{num}.  ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = AZUL
    r = p.add_run(titulo)
    r.bold = True; r.font.size = Pt(10)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(4)
    p2.paragraph_format.left_indent  = Cm(1.2)
    r2 = p2.add_run(desc)
    r2.font.size = Pt(9)
    r2.font.color.rgb = CINZA

caixa("Não feche a janela do terminal enquanto estiver usando o sistema. "
      "Ao fechar o terminal, o servidor é encerrado e a aplicação para de funcionar.",
      bold_prefix="Atenção:")

# ── 4. GUIA DE USO ─────────────────────────────────────────────────────────────
heading("4. Guia de Uso — Passo a Passo")
hr_line()
para("O sistema é organizado em 6 etapas visíveis no painel lateral. "
     "As etapas 1 a 4 formam o fluxo principal de anonimização. "
     "A etapa 5 é para reversão e a 6 para PDFs.", size=10, space_after=6)

etapas = [
    ("Etapa", "Nome", "Função", True),
    ("1", "Carregar Arquivo",   "Selecionar o arquivo de dados a processar", False),
    ("2", "Analisar Estrutura", "Visualizar colunas e campos detectados", False),
    ("3", "Selecionar Campos",  "Escolher quais dados serão anonimizados", False),
    ("4", "Anonimizar Dados",   "Gerar arquivo anonimizado e chave JSON", False),
    ("5", "Reverter Tokens",    "Restaurar dados originais com a chave JSON", False),
    ("6", "Processar PDF / OCR","Anonimizar documentos PDF com OCR", False),
]
t_et = doc.add_table(rows=len(etapas), cols=3)
t_et.alignment = WD_TABLE_ALIGNMENT.LEFT
w_et = [Cm(1.5), Cm(5), Cm(9.5)]
for i, row_data in enumerate(etapas):
    c1, c2, c3, hdr = row_data
    for j, (cell, txt, w) in enumerate(zip(t_et.rows[i].cells, [c1, c2, c3], w_et)):
        cell.width = w
        set_cell_border(cell)
        if hdr:
            set_cell_bg(cell, "2563EB")
        elif i % 2 == 0:
            set_cell_bg(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        if j == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.bold = hdr; r.font.size = Pt(9)
        r.font.color.rgb = BRANCO if hdr else RGBColor(0x1E, 0x29, 0x3B)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# 4.1
subheading("4.1  Carregar Arquivo")
hr_line("CBD5E1")
para("Na tela inicial, arraste o arquivo para a área indicada ou clique para selecionar pelo explorador. "
     "Após o carregamento, o sistema avança automaticamente para a etapa 2.", size=10, space_after=4)
caixa_info("Formatos aceitos: .xlsx, .csv, .docx e .txt. Para PDF, utilize a etapa 6.")

# 4.2
subheading("4.2  Analisar Estrutura")
hr_line("CBD5E1")
para("O sistema exibe a estrutura interna do arquivo: abas (Excel), colunas, tipos inferidos e amostras. "
     "Esta etapa é apenas de leitura. Clique em Avançar para continuar.", size=10, space_after=6)

# 4.3
subheading("4.3  Selecionar Campos")
hr_line("CBD5E1")
para("A tabela lista todos os campos detectados com classificação automática de risco:", size=10, space_after=4)

risco = [
    ("Classificação", "Cor", "Exemplos", True),
    ("CRÍTICO",  "Vermelho", "CNPJ, CPF", False),
    ("ALTO",     "Laranja",  "Razão Social, Nome, Empresa", False),
    ("MÉDIO",    "Amarelo",  "E-mail, Endereço, Município", False),
    ("BAIXO",    "Cinza",    "Campos genéricos", False),
]
t_r = doc.add_table(rows=len(risco), cols=3)
t_r.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, (c1, c2, c3, hdr) in enumerate(risco):
    for j, (cell, txt, w) in enumerate(zip(t_r.rows[i].cells, [c1, c2, c3],
                                            [Cm(3), Cm(3), Cm(10)])):
        cell.width = w
        set_cell_border(cell)
        set_cell_bg(cell, "0F172A" if hdr else ("F8FAFC" if i % 2 == 0 else "FFFFFF"))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(txt)
        r.bold = hdr; r.font.size = Pt(9)
        r.font.color.rgb = BRANCO if hdr else RGBColor(0x1E, 0x29, 0x3B)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
para("Use os botões de alternância na última coluna para ativar/desativar cada campo. "
     "Campos CRÍTICO e ALTO ficam ativados por padrão. Clique em Anonimizar para prosseguir.",
     size=10, space_after=6)

# 4.4
page_break()
subheading("4.4  Anonimizar Dados")
hr_line("CBD5E1")
para("Após a tokenização, dois arquivos ficam disponíveis para download:", size=10, space_after=4)

dl = [
    ("Arquivo", "Descrição", True),
    ("Salvar Arquivo",
     "Documento anonimizado no mesmo formato original. Dados sensíveis substituídos por tokens [TKN-XXXXXXXX].", False),
    ("Salvar Chave JSON",
     "Dicionário de reversão. Guarde com segurança — sem ele não é possível recuperar os dados originais.", False),
]
t_dl = doc.add_table(rows=len(dl), cols=2)
t_dl.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, (c1, c2, hdr) in enumerate(dl):
    for j, (cell, txt, w) in enumerate(zip(t_dl.rows[i].cells, [c1, c2], [Cm(4), Cm(12)])):
        cell.width = w
        set_cell_border(cell)
        set_cell_bg(cell, "0F172A" if hdr else ("F8FAFC" if i % 2 == 0 else "FFFFFF"))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(txt)
        r.bold = hdr; r.font.size = Pt(9)
        r.font.color.rgb = BRANCO if hdr else RGBColor(0x1E, 0x29, 0x3B)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

doc.add_paragraph().paragraph_format.space_after = Pt(2)
caixa("Salve sempre os dois arquivos. A chave JSON é indispensável para reverter a anonimização.",
      bold_prefix="Importante:")

# 4.5
subheading("4.5  Reverter Tokens")
hr_line("CBD5E1")
para("Para restaurar os dados originais de um arquivo anonimizado:", size=10, space_after=4)
for n, txt in [
    ("1", "Acesse a etapa 5. Reverter Tokens no menu lateral."),
    ("2", "Carregue o arquivo anonimizado no campo 'Arquivo Anonimizado'."),
    ("3", "Carregue o arquivo .json no campo 'Chave de Reversão'."),
    ("4", "Clique em Reverter Dados para baixar o arquivo restaurado."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.5)
    r = p.add_run(f"{n}.  ")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = AZUL
    r = p.add_run(txt)
    r.font.size = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
caixa_info("A reversão é determinística: os mesmos tokens sempre produzem os mesmos valores originais "
           "com a chave JSON correta.")

# 4.6
subheading("4.6  Processar PDF / OCR")
hr_line("CBD5E1")
para("Esta etapa permite anonimizar documentos PDF, incluindo PDFs digitalizados via OCR (Tesseract).", size=10, space_after=4)
for bullet, txt in [
    ("PDF digital:", "Processado diretamente pelo PDF.js."),
    ("PDF escaneado:", "O OCR extrai o texto das imagens antes da anonimização."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.5)
    r = p.add_run(f"• {bullet}  ")
    r.bold = True; r.font.size = Pt(10)
    r = p.add_run(txt)
    r.font.size = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
caixa("Esta funcionalidade exige acesso via servidor local (http://localhost:8000). "
      "Não funciona ao abrir o arquivo diretamente pelo Windows Explorer (file://).",
      bold_prefix="Requisito:")

# ── 5. SEGURANÇA ───────────────────────────────────────────────────────────────
heading("5. Segurança e Privacidade")
hr_line()

seg = [
    ("Característica", "Detalhes", True),
    ("Processamento local",
     "Todo o processamento ocorre no computador local. Nenhum dado é enviado a servidores externos.", False),
    ("Sem dependência de rede",
     "O sistema funciona completamente offline após a instalação.", False),
    ("Tokenização determinística",
     "O mesmo valor sempre gera o mesmo token, permitindo consistência entre arquivos.", False),
    ("Chave reversível",
     "A reversão só é possível com a chave JSON gerada no momento da anonimização.", False),
    ("Sem armazenamento persistente",
     "Dados processados existem apenas na memória RAM. Ao fechar o navegador, são descartados.", False),
]
t_seg = doc.add_table(rows=len(seg), cols=2)
t_seg.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, (c1, c2, hdr) in enumerate(seg):
    for j, (cell, txt, w) in enumerate(zip(t_seg.rows[i].cells, [c1, c2], [Cm(5), Cm(11)])):
        cell.width = w
        set_cell_border(cell, "D1FAE5")
        set_cell_bg(cell, "064E3B" if hdr else ("F0FDF4" if i % 2 == 0 else "FFFFFF"))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(txt)
        r.bold = hdr; r.font.size = Pt(9)
        r.font.color.rgb = BRANCO if hdr else RGBColor(0x1E, 0x29, 0x3B)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ── Rodapé ─────────────────────────────────────────────────────────────────────
from docx.oxml import OxmlElement as OE2
section = doc.sections[0]
footer  = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Anônimo v1.0  ·  Manual do Utilizador  ·  Coordenadoria de Inteligência Fiscal")
r.font.size = Pt(8)
r.font.color.rgb = CINZA

# ── Salvar ─────────────────────────────────────────────────────────────────────
doc.save(OUTPUT)
print(f"Documento gerado: {OUTPUT}")
